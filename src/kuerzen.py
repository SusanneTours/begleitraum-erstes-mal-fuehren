"""Erzeugt die gekürzte Fassung (so wie das Training in Köln gelaufen ist) aus der Vollversion."""
import sys
from pptx import Presentation
RAUS = [6, 8, 12, 14, 15, 18, 19, 20, 35, 36, 37, 38, 39, 42, 44, 45, 47, 48, 49, 50, 51, 52, 54, 59, 60, 61, 62, 63, 64, 65, 66, 67, 68, 69]
src, out = sys.argv[1], sys.argv[2]
p = Presentation(src)
ids = list(p.slides._sldIdLst)
for n in sorted(RAUS, reverse=True):
    sid = ids[n - 1]
    p.part.drop_rel(sid.rId)
    p.slides._sldIdLst.remove(sid)
p.save(out)
print('Folien:', len(p.slides))
